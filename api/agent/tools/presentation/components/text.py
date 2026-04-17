"""Text component — fills any placeholder with text or bullet list."""

import re
from typing import Optional

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

from api.agent.tools.presentation.theme import TemplateTheme, parse_md_segments, set_run_font


def fill_text(
    placeholder,
    spec: dict,
    theme: TemplateTheme,
) -> None:
    """Fill a placeholder with text or bullet content.

    Args:
        placeholder: A pptx placeholder shape with a text_frame.
        spec: Component spec with 'text' and/or 'bullets', plus optional 'style'.
        theme: The active template theme.
    """
    tf = placeholder.text_frame
    tf.word_wrap = True
    style = spec.get("style", "body")

    # Clear existing placeholder text
    for para in tf.paragraphs:
        para.text = ""

    bullets = spec.get("bullets", [])
    text = spec.get("text", "")

    if bullets:
        _render_bullets(tf, bullets, style, theme)
    elif text:
        _render_text(tf, text, style, theme)


def _render_text(
    tf,
    text: str,
    style: str,
    theme: TemplateTheme,
) -> None:
    """Render a single text block with markdown bold support."""
    size_pt = {"heading": 28, "subtitle": 16, "body": 14}.get(style, 14)
    color = theme.fg
    bold_color = theme.accent

    para = tf.paragraphs[0]
    for segment, is_bold in parse_md_segments(text):
        run = para.add_run()
        run.text = segment
        set_run_font(
            run, size_pt, theme,
            bold=is_bold,
            color=bold_color if is_bold else color,
        )


def _render_bullets(
    tf,
    bullets: list[str],
    style: str,
    theme: TemplateTheme,
) -> None:
    """Render a bullet list with markdown bold support."""
    size_pt = {"heading": 16, "subtitle": 14, "body": 13}.get(style, 13)

    for i, bullet in enumerate(bullets):
        # Strip leading bullet markers (-, *, •) but preserve markdown **bold**
        clean = re.sub(r"^[\-\u2022]\s*", "", str(bullet).strip())
        # Only strip a single leading * if it's NOT part of ** (markdown bold)
        if clean.startswith("*") and not clean.startswith("**"):
            clean = clean[1:].lstrip()

        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = PP_ALIGN.LEFT
        if i > 0:
            para.space_before = Pt(10)

        # Bullet marker
        marker = para.add_run()
        marker.text = "\u25b8  "  # triangle bullet
        set_run_font(marker, size_pt, theme, bold=True, color=theme.accent)

        # Content with bold support
        for segment, is_bold in parse_md_segments(clean):
            run = para.add_run()
            run.text = segment
            set_run_font(
                run, size_pt, theme,
                bold=is_bold,
                color=theme.accent if is_bold else theme.fg,
            )
